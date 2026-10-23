import openpyxl#导入openpyxl模块：专门用来读写xlsx文件的工具包，要用它得先导入
wb = openpyxl.load_workbook(r"C:\Users\Administrator\Desktop\AI交叉岗位情报站.xlsx")#打开整个Excel文件，装进变量wb里（wb=工作簿，一个文件=一本账本）
ws = wb["岗位收集"]#从工作簿里按名字翻到"岗位收集"这一页表（ws=工作表，wb是本子，ws是其中一页）
# row = ws[2]#取第2行：拿到的是这一行12个格子打包成的一组，row[0]是A列、row[1]是B列……程序员从0开始数
# print(row[1].value, row[8].value, "-", row[9].value, "k/月")#打印B2(岗位名)、I2(薪资下限)、J2(薪资上限)，.value才是取出格子里装的内容
# row = ws[3]#换成第3行（我自己加的改动）
# row[8].value = 15#把I3格子的值改成15（注意：只是改了内存里的副本，Excel文件本身没变）
# row[9].value= 25#把J3格子的值改成25（同上，没存盘）
# print(row[1].value, row[8].value, "-", row[9].value, "k/月")#再打印一遍：输出的是改过之后的数
# wb.save(r"C:\Users\Administrator\Desktop\AI交叉岗位情报站.xlsx")#'r'-raw string,字符就只是字符，如\n.只是防转义的 "保险"，它既不是 "读取" 的意思，也不是 "读取文件夹" 的方式。
for row in ws.iter_rows():#外层循环：一行一行翻这张表（表是二维的，行一个方向列一个方向，一个循环只能管一个方向）
    for cell in row:#内层循环：在这一行里一个格子一个格子地过
        print(cell.value,end =' ')#打印格子内容，end=' '让同一行的数据用空格隔开而不是换行
    print()#一整行打完，顶一个换行，让下一行数据另起一行